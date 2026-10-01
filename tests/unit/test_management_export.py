import csv
import io
from datetime import UTC, datetime
from uuid import UUID, uuid4
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from easyaudit_next.management.export import (
    EXPORT_COLUMNS,
    ExportFormat,
    protect_text,
    render_export,
)
from easyaudit_next.management.schemas import (
    ActionLifecycleCounts,
    DeadlineBucket,
    FindingLifecycleCounts,
    ManagementCaseExportSnapshot,
    ManagementCaseFilters,
    ManagementCaseSummary,
    ManagementDeadlineFilter,
)
from easyaudit_next.review_core.domain.models import ReviewCaseLifecycle

SHANGHAI = ZoneInfo("Asia/Shanghai")
GENERATED_AT = datetime(2026, 10, 2, 1, 30, 5, tzinfo=UTC)  # 09:30 in Shanghai
AS_OF = datetime(2026, 10, 2, 1, 29, 59, 123456, tzinfo=UTC)

INJECTION_TITLES = [
    "=1+1",
    "+cmd",
    "-2+3",
    "@SUM(A1)",
    "\t=tab",
    "\r=return",
]


def _summary(
    title: str,
    *,
    planned_end_at: datetime | None = None,
    review_plan_id: UUID | None = None,
) -> ManagementCaseSummary:
    return ManagementCaseSummary(
        id=uuid4(),
        review_plan_id=review_plan_id,
        title=title,
        scenario_key="process_review",
        scenario_version=1,
        lifecycle=ReviewCaseLifecycle.IN_PROGRESS,
        planned_start_at=None,
        planned_end_at=planned_end_at,
        deadline_bucket=DeadlineBucket.NONE,
        findings=FindingLifecycleCounts(
            total=3, open=1, rectifying=1, verifying=0, closed=1, voided=0
        ),
        actions=ActionLifecycleCounts(
            total=2, todo=1, in_progress=0, done=1, cancelled=0, overdue=1, due_soon=0
        ),
    )


def _snapshot(
    titles: list[str],
    *,
    organization_name: str = "Acme",
    filters: ManagementCaseFilters | None = None,
) -> ManagementCaseExportSnapshot:
    return ManagementCaseExportSnapshot(
        as_of=AS_OF,
        organization_name=organization_name,
        filters=filters or ManagementCaseFilters(),
        items=tuple(_summary(title) for title in titles),
    )


def _render(snapshot: ManagementCaseExportSnapshot, fmt: ExportFormat):  # type: ignore[no-untyped-def]
    return render_export(snapshot, fmt, generated_at=GENERATED_AT, tz=SHANGHAI)


def _read_csv(content: bytes) -> tuple[list[list[str]], list[list[str]]]:
    assert content.startswith(b"\xef\xbb\xbf"), "CSV must start with a UTF-8 BOM"
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"), newline="")))
    blank = rows.index([])
    return rows[:blank], rows[blank + 1 :]


@pytest.mark.parametrize("title", INJECTION_TITLES)
def test_csv_neutralizes_formula_prefixes_in_text_fields(title: str) -> None:
    document = _render(_snapshot([title], organization_name=title), ExportFormat.CSV)

    metadata, table = _read_csv(document.content)
    header, row = table
    assert header == list(EXPORT_COLUMNS)
    assert row[header.index("title")] == "'" + title
    assert dict(metadata)["organization_name"] == "'" + title


@pytest.mark.parametrize("title", INJECTION_TITLES)
def test_xlsx_stores_formula_prefixes_as_strings_not_formulas(title: str) -> None:
    document = _render(_snapshot([title], organization_name=title), ExportFormat.XLSX)

    workbook = load_workbook(io.BytesIO(document.content))
    # XML normalizes a bare carriage return to a line feed; the `'` guard is what matters.
    expected = "'" + title.replace("\r", "\n")
    cases = workbook["review_cases"]
    title_cell = cases.cell(row=2, column=EXPORT_COLUMNS.index("title") + 1)
    assert title_cell.data_type == "s"
    assert title_cell.value == expected
    organization_cell = workbook["metadata"]["B4"]
    assert organization_cell.data_type == "s"
    assert organization_cell.value == expected
    # No cell anywhere is a formula.
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            assert all(cell.data_type != "f" for cell in row)


def test_protect_text_leaves_safe_text_alone_and_replaces_xml_illegal_characters() -> None:
    assert protect_text("plain 标题") == "plain 标题"
    assert protect_text("a=b") == "a=b"
    assert protect_text("") == ""
    assert protect_text("x\x01y") == "x�y"


TRICKY_TITLE = '审查 "第一阶段", 含 emoji 🚀\n第二行,逗号'


def test_csv_round_trips_unicode_newlines_commas_and_quotes() -> None:
    document = _render(_snapshot([TRICKY_TITLE]), ExportFormat.CSV)

    _, table = _read_csv(document.content)
    assert table[1][table[0].index("title")] == TRICKY_TITLE


def test_xlsx_round_trips_unicode_newlines_commas_and_quotes() -> None:
    document = _render(_snapshot([TRICKY_TITLE]), ExportFormat.XLSX)

    cases = load_workbook(io.BytesIO(document.content))["review_cases"]
    assert cases.cell(row=2, column=EXPORT_COLUMNS.index("title") + 1).value == TRICKY_TITLE


def test_counters_stay_numeric_in_both_formats() -> None:
    snapshot = _snapshot(["Case"])
    cases = load_workbook(io.BytesIO(_render(snapshot, ExportFormat.XLSX).content))[
        "review_cases"
    ]
    column = EXPORT_COLUMNS.index("findings_total") + 1
    assert cases.cell(row=2, column=column).value == 3
    assert cases.cell(row=2, column=column).data_type == "n"


def test_times_use_the_export_timezone_with_offset_and_metadata_names_it() -> None:
    planned_end = datetime(2026, 10, 1, 16, 0, 0, 654321, tzinfo=UTC)
    snapshot = ManagementCaseExportSnapshot(
        as_of=AS_OF,
        organization_name="Acme",
        filters=ManagementCaseFilters(
            lifecycle=ReviewCaseLifecycle.IN_PROGRESS,
            deadline_status=ManagementDeadlineFilter.OVERDUE,
        ),
        items=(_summary("Case", planned_end_at=planned_end),),
    )

    metadata, table = _read_csv(_render(snapshot, ExportFormat.CSV).content)

    header, row = table
    assert row[header.index("planned_end_at")] == "2026-10-02T00:00:00.654321+08:00"
    assert row[header.index("planned_start_at")] == ""
    assert dict(metadata) == {
        "generated_at": "2026-10-02T09:30:05+08:00",
        "as_of": "2026-10-02T09:29:59.123456+08:00",
        "timezone": "Asia/Shanghai",
        "organization_name": "Acme",
        "filter_review_plan_id": "",
        "filter_lifecycle": "in_progress",
        "filter_deadline_status": "overdue",
        "row_count": "1",
    }


def test_xlsx_has_review_cases_and_metadata_sheets() -> None:
    workbook = load_workbook(io.BytesIO(_render(_snapshot(["A", "B"]), ExportFormat.XLSX).content))

    assert workbook.sheetnames == ["review_cases", "metadata"]
    metadata = {row[0].value: row[1].value for row in workbook["metadata"].iter_rows()}
    assert metadata["timezone"] == "Asia/Shanghai"
    assert metadata["row_count"] == 2
    assert metadata["generated_at"] == "2026-10-02T09:30:05+08:00"
    assert [cell.value for cell in workbook["review_cases"][1]] == list(EXPORT_COLUMNS)


def test_other_timezone_is_honored() -> None:
    tokyo = ZoneInfo("Asia/Tokyo")
    document = render_export(
        _snapshot(["A"]), ExportFormat.CSV, generated_at=GENERATED_AT, tz=tokyo
    )
    metadata, _ = _read_csv(document.content)
    assert dict(metadata)["generated_at"] == "2026-10-02T10:30:05+09:00"
    assert document.filename == "review-cases-20261002T1030+0900.csv"


def test_filename_and_headers_are_safe_and_timestamped() -> None:
    csv_document = _render(_snapshot([]), ExportFormat.CSV)
    xlsx_document = _render(_snapshot([]), ExportFormat.XLSX)

    assert csv_document.filename == "review-cases-20261002T0930+0800.csv"
    assert xlsx_document.filename == "review-cases-20261002T0930+0800.xlsx"
    assert csv_document.media_type == "text/csv; charset=utf-8"
    assert xlsx_document.media_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert csv_document.content_disposition == (
        'attachment; filename="review-cases-20261002T0930+0800.csv"; '
        "filename*=UTF-8''review-cases-20261002T0930%2B0800.csv"
    )


def test_empty_export_still_has_metadata_and_header() -> None:
    metadata, table = _read_csv(_render(_snapshot([]), ExportFormat.CSV).content)

    assert dict(metadata)["row_count"] == "0"
    assert table == [list(EXPORT_COLUMNS)]


@pytest.mark.parametrize("char", ["\ufffe", "\uffff", "\x01", "\x1f", "\x0b"])
def test_xml_illegal_characters_are_replaced_and_both_formats_stay_readable(char: str) -> None:
    title = f"a{char}b"
    snapshot = _snapshot([title], organization_name=title)

    csv_metadata, csv_table = _read_csv(_render(snapshot, ExportFormat.CSV).content)
    assert csv_table[1][csv_table[0].index("title")] == "a\ufffdb"
    assert dict(csv_metadata)["organization_name"] == "a\ufffdb"

    workbook = load_workbook(io.BytesIO(_render(snapshot, ExportFormat.XLSX).content))
    cell = workbook["review_cases"].cell(row=2, column=EXPORT_COLUMNS.index("title") + 1)
    assert cell.value == "a\ufffdb"
    assert workbook["metadata"]["B4"].value == "a\ufffdb"


def test_legal_boundary_characters_are_preserved() -> None:
    for char in ("\t", "\n", "\ud7ff", "\ue000", "\ufffd", "\U00010000", "\U0010ffff"):
        assert protect_text(f"a{char}b") == f"a{char}b"
