"""Management export: same rows as the UI JSON, authorization, limits, headers, logging."""

import csv
import io
import logging
import os
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from typing import Any
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, create_engine
from sqlalchemy.orm import Session

from easyaudit_next.api.dependencies import get_current_identity, get_database_session
from easyaudit_next.infrastructure.observability import APP_LOGGER, JsonFormatter
from easyaudit_next.main import create_app
from easyaudit_next.management import api as management_api
from easyaudit_next.management.export import EXPORT_COLUMNS
from easyaudit_next.platform.domain.ids import OrganizationId, UserId
from easyaudit_next.platform.persistence.models import OrganizationRecord, UserRecord
from easyaudit_next.platform.settings import get_settings
from easyaudit_next.review_core.persistence.models import (
    ActionItemRecord,
    CaseMemberRecord,
    FindingRecord,
    ReviewCaseRecord,
    ReviewPlanRecord,
    ScenarioRecord,
    ScenarioVersionRecord,
)
from tests.integration.test_management_api import _identity

EXPORT_URL = "/api/v1/management/review-cases/export"
LIST_URL = "/api/v1/management/review-cases"
CSV_TYPE = "text/csv; charset=utf-8"


@pytest.fixture(scope="module")
def postgres_engine() -> Iterator[Engine]:
    if os.getenv("EASYAUDIT_RUN_POSTGRES_TESTS") != "1":
        pytest.skip("PostgreSQL integration tests are opt-in outside CI")
    engine = create_engine(os.environ["DATABASE_URL"], pool_pre_ping=True)
    yield engine
    engine.dispose()


@dataclass(frozen=True)
class Seed:
    organization_id: OrganizationId
    lead_id: UserId
    second_lead_id: UserId
    observer_id: UserId
    plan_id: UUID
    visible_case_ids: tuple[UUID, ...]


CASE_TITLES = [
    "=1+1",
    "+cmd|' /C calc'!A0",
    "-2+3",
    "@SUM(A1)",
    "\t制表符开头",
    '中文标题 "引号", 逗号 🚀\n第二行',
    "Plain English",
]


def _seed(engine: Engine) -> Seed:
    now = datetime.now(UTC).replace(microsecond=123456)
    organization_id = OrganizationId(uuid4())
    other_org_id = OrganizationId(uuid4())
    lead_id, second_lead_id, observer_id, outsider_id = (UserId(uuid4()) for _ in range(4))
    plan_id = uuid4()
    case_ids = [uuid4() for _ in CASE_TITLES]
    hidden_case_id = uuid4()
    observer_case_id = uuid4()

    with Session(engine) as session, session.begin():
        session.add_all(
            [
                OrganizationRecord(id=organization_id, name="=审计中心, \"一部\""),
                OrganizationRecord(id=other_org_id, name="Other org"),
            ]
        )
        session.flush()
        session.add_all(
            [
                UserRecord(
                    id=user_id,
                    organization_id=organization_id,
                    display_name=name,
                    platform_role="ordinary_user",
                )
                for user_id, name in (
                    (lead_id, "Lead"),
                    (second_lead_id, "Second lead"),
                    (observer_id, "Observer"),
                )
            ]
        )
        session.add(
            UserRecord(
                id=outsider_id,
                organization_id=other_org_id,
                display_name="Outsider",
                platform_role="ordinary_user",
            )
        )
        session.flush()
        scenario_ids: dict[UUID, UUID] = {}
        version_ids: dict[UUID, UUID] = {}
        for org in (organization_id, other_org_id):
            scenario_ids[org] = uuid4()
            version_ids[org] = uuid4()
            session.add(
                ScenarioRecord(
                    id=scenario_ids[org],
                    organization_id=org,
                    key="process_review",
                    name="Process Review",
                )
            )
            session.flush()
            session.add(
                ScenarioVersionRecord(
                    id=version_ids[org],
                    scenario_id=scenario_ids[org],
                    organization_id=org,
                    version=1,
                    published_at=now,
                )
            )
        session.flush()
        session.add(
            ReviewPlanRecord(
                id=plan_id,
                organization_id=organization_id,
                title="Annual plan",
                planned_start_at=None,
                planned_end_at=None,
                created_by=lead_id,
                created_at=now,
            )
        )
        session.flush()

        def case(
            case_id: UUID, org: UUID, title: str, index: int, **extra: Any
        ) -> ReviewCaseRecord:
            return ReviewCaseRecord(
                id=case_id,
                organization_id=org,
                plan_id=extra.pop("plan_id", None),
                scenario_version_id=version_ids[org],
                title=title,
                lifecycle=extra.pop("lifecycle", "in_progress"),
                planned_start_at=now - timedelta(days=30),
                planned_end_at=now + timedelta(days=index - 3, hours=index),
                started_at=now - timedelta(days=30),
                fieldwork_completed_at=None,
                closed_at=None,
                scenario_data_json={},
                created_by=outsider_id if org == other_org_id else lead_id,
                created_at=now - timedelta(days=31),
            )

        cases = [
            case(
                case_id,
                organization_id,
                title,
                index,
                plan_id=plan_id if index % 2 == 0 else None,
                lifecycle="in_progress" if index != 5 else "scheduled",
            )
            for index, (case_id, title) in enumerate(zip(case_ids, CASE_TITLES, strict=True))
        ]
        cases.append(case(hidden_case_id, organization_id, "Only second lead", 1))
        cases.append(case(observer_case_id, organization_id, "Observer only", 1))
        cases.append(case(uuid4(), other_org_id, "Other org case", 1))
        session.add_all(cases)
        session.flush()
        members = [
            CaseMemberRecord(
                organization_id=organization_id,
                case_id=case_id,
                user_id=lead_id,
                role_key="lead",
                joined_at=now,
            )
            for case_id in case_ids
        ]
        members.append(
            CaseMemberRecord(
                organization_id=organization_id,
                case_id=hidden_case_id,
                user_id=second_lead_id,
                role_key="lead",
                joined_at=now,
            )
        )
        members.append(
            CaseMemberRecord(
                organization_id=organization_id,
                case_id=observer_case_id,
                user_id=observer_id,
                role_key="observer",
                joined_at=now,
            )
        )
        session.add_all(members)
        session.flush()
        for index, case_id in enumerate(case_ids[:5]):
            for finding_index in range(index % 3 + 1):
                finding_id = uuid4()
                session.add(
                    FindingRecord(
                        id=finding_id,
                        organization_id=organization_id,
                        case_id=case_id,
                        title=f"Finding {finding_index}",
                        description=None,
                        severity="high",
                        lifecycle=("open", "rectifying", "closed")[finding_index % 3],
                        raised_by=lead_id,
                        raised_at=now - timedelta(hours=finding_index + 1),
                        scenario_data_json={},
                    )
                )
                session.flush()
                session.add(
                    ActionItemRecord(
                        id=uuid4(),
                        organization_id=organization_id,
                        finding_id=finding_id,
                        title="Action",
                        lifecycle=("todo", "in_progress", "done")[finding_index % 3],
                        due_at=now + timedelta(days=finding_index * 4 - 2),
                        completed_at=None,
                    )
                )
    return Seed(
        organization_id,
        lead_id,
        second_lead_id,
        observer_id,
        plan_id,
        tuple(case_ids),
    )


def _client(engine: Engine, organization_id: OrganizationId, user_id: UserId) -> TestClient:
    app = create_app()
    identity = _identity(organization_id, user_id)
    app.dependency_overrides[get_current_identity] = lambda: identity

    def database_session() -> Iterator[Session]:
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_database_session] = database_session
    return TestClient(app)


TIME_COLUMNS = (
    EXPORT_COLUMNS.index("planned_start_at"),
    EXPORT_COLUMNS.index("planned_end_at"),
)


def _instant(value: str | None) -> datetime | str:
    """Time cell as an exact instant; the export must carry the +08:00 offset."""
    if not value:
        return ""
    parsed = datetime.fromisoformat(value)
    assert parsed.utcoffset() == timedelta(hours=8), value
    return parsed


def _with_instants(row: list[Any]) -> list[Any]:
    return [
        _instant(value) if index in TIME_COLUMNS else value for index, value in enumerate(row)
    ]


def _guard(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def _json_instant(value: str | None) -> datetime | str:
    # The UI JSON is the reference: compare the exact instant, whatever offset it is written in.
    return "" if value is None else datetime.fromisoformat(value)


def _row_from_json(item: dict[str, Any]) -> list[str | int]:
    """Independent restatement of the export columns from one UI JSON item."""
    findings, actions = item["findings"], item["actions"]
    return [
        item["id"],
        item["review_plan_id"] or "",
        _guard(item["title"]),
        item["scenario_key"],
        item["scenario_version"],
        item["lifecycle"],
        _json_instant(item["planned_start_at"]),
        _json_instant(item["planned_end_at"]),
        item["deadline_bucket"],
        findings["total"],
        findings["open"],
        findings["rectifying"],
        findings["verifying"],
        findings["closed"],
        findings["voided"],
        actions["total"],
        actions["todo"],
        actions["in_progress"],
        actions["done"],
        actions["cancelled"],
        actions["overdue"],
        actions["due_soon"],
    ]


def _json_rows(client: TestClient, query: str = "") -> list[list[str | int]]:
    items: list[dict[str, Any]] = []
    offset = 0
    while True:
        separator = "&" if query else ""
        response = client.get(f"{LIST_URL}?{query}{separator}limit=2&offset={offset}")
        assert response.status_code == 200
        page = response.json()
        items.extend(page["items"])
        offset += 2
        if offset >= page["total"]:
            return [_row_from_json(item) for item in items]


def _csv_rows(client: TestClient, query: str = "") -> list[list[str | int]]:
    response = client.get(f"{EXPORT_URL}?format=csv{'&' + query if query else ''}")
    assert response.status_code == 200
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"), newline="")))
    header, *data = rows[rows.index([]) + 1 :]
    assert header == list(EXPORT_COLUMNS)
    counters = ("findings_", "actions_")
    numeric = {i for i, name in enumerate(EXPORT_COLUMNS) if name.startswith(counters)}
    numeric.add(EXPORT_COLUMNS.index("scenario_version"))
    return [
        _with_instants(
            [int(value) if index in numeric else value for index, value in enumerate(row)]
        )
        for row in data
    ]


def _xlsx_rows(client: TestClient, query: str = "") -> list[list[str | int]]:
    response = client.get(f"{EXPORT_URL}?format=xlsx{'&' + query if query else ''}")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content))["review_cases"]
    header, *data = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert header == list(EXPORT_COLUMNS)
    return [_with_instants(["" if value is None else value for value in row]) for row in data]


@pytest.mark.parametrize(
    "query",
    ["", "lifecycle=scheduled", "deadline_status=overdue", "deadline_status=due_soon", "plan"],
)
def test_json_csv_and_xlsx_rows_are_identical(postgres_engine: Engine, query: str) -> None:
    seed = _seed(postgres_engine)
    client = _client(postgres_engine, seed.organization_id, seed.lead_id)
    if query == "plan":
        query = f"review_plan_id={seed.plan_id}"

    expected = _json_rows(client, query)
    assert expected
    if query == "":
        assert len(expected) == len(CASE_TITLES)
    assert _csv_rows(client, query) == expected
    assert _xlsx_rows(client, query) == expected


def test_export_rows_match_what_each_user_sees_in_the_ui(postgres_engine: Engine) -> None:
    seed = _seed(postgres_engine)

    lead = _client(postgres_engine, seed.organization_id, seed.lead_id)
    second = _client(postgres_engine, seed.organization_id, seed.second_lead_id)
    observer = _client(postgres_engine, seed.organization_id, seed.observer_id)

    lead_rows = _csv_rows(lead)
    assert {row[0] for row in lead_rows} == {str(case_id) for case_id in seed.visible_case_ids}
    assert lead_rows == _json_rows(lead)

    second_rows = _csv_rows(second)
    assert [row[2] for row in second_rows] == ["Only second lead"]
    assert second_rows == _json_rows(second) == _xlsx_rows(second)

    # Observer holds no manage permission: nothing to export, and nothing leaks via metadata.
    assert _csv_rows(observer) == []
    assert _xlsx_rows(observer) == []


def test_export_is_ordered_like_the_ui_list(postgres_engine: Engine) -> None:
    seed = _seed(postgres_engine)
    client = _client(postgres_engine, seed.organization_id, seed.lead_id)

    ui_ids = [
        item["id"] for item in client.get(f"{LIST_URL}?limit=100").json()["items"]
    ]
    assert [row[0] for row in _csv_rows(client)] == ui_ids


def test_export_response_headers(postgres_engine: Engine) -> None:
    seed = _seed(postgres_engine)
    client = _client(postgres_engine, seed.organization_id, seed.lead_id)

    csv_response = client.get(f"{EXPORT_URL}?format=csv")
    xlsx_response = client.get(f"{EXPORT_URL}?format=xlsx")

    for response, extension in ((csv_response, "csv"), (xlsx_response, "xlsx")):
        assert response.status_code == 200
        assert response.headers["x-content-type-options"] == "nosniff"
        assert response.headers["cache-control"] == "no-store"
        disposition = response.headers["content-disposition"]
        assert disposition.startswith('attachment; filename="review-cases-')
        assert disposition.split('filename="')[1].split('"')[0].endswith(f"+0800.{extension}")
    assert csv_response.headers["content-type"] == CSV_TYPE
    assert xlsx_response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_export_requires_an_explicit_valid_format(postgres_engine: Engine) -> None:
    seed = _seed(postgres_engine)
    client = _client(postgres_engine, seed.organization_id, seed.lead_id)

    missing = client.get(EXPORT_URL)
    assert missing.status_code == 422
    assert missing.json()["errors"] == [{"field": "format", "code": "required", "params": {}}]
    invalid = client.get(f"{EXPORT_URL}?format=pdf")
    assert invalid.status_code == 422
    assert invalid.json()["errors"][0]["field"] == "format"


def test_export_metadata_carries_organization_filters_and_timezone(
    postgres_engine: Engine,
) -> None:
    seed = _seed(postgres_engine)
    client = _client(postgres_engine, seed.organization_id, seed.lead_id)

    query = f"format=csv&lifecycle=scheduled&review_plan_id={seed.plan_id}"
    response = client.get(f"{EXPORT_URL}?{query}")

    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"), newline="")))
    metadata = dict(rows[: rows.index([])])
    assert metadata["organization_name"] == "'=审计中心, \"一部\""
    assert metadata["timezone"] == "Asia/Shanghai"
    assert metadata["filter_lifecycle"] == "scheduled"
    assert metadata["filter_review_plan_id"] == str(seed.plan_id)
    assert metadata["filter_deadline_status"] == "all"
    assert metadata["generated_at"].endswith("+08:00")
    assert metadata["as_of"].endswith("+08:00")
    assert "." in metadata["as_of"]  # sub-second precision is kept


def _with_max_rows(monkeypatch: pytest.MonkeyPatch, max_rows: int) -> None:
    settings = get_settings().model_copy(update={"export_max_rows": max_rows})
    monkeypatch.setattr(management_api, "get_settings", lambda: settings)


@pytest.mark.parametrize("export_format", ["csv", "xlsx"])
def test_export_over_the_row_limit_fails_whole_without_partial_output(
    postgres_engine: Engine, monkeypatch: pytest.MonkeyPatch, export_format: str
) -> None:
    seed = _seed(postgres_engine)
    client = _client(postgres_engine, seed.organization_id, seed.lead_id)
    _with_max_rows(monkeypatch, len(CASE_TITLES) - 1)

    response = client.get(f"{EXPORT_URL}?format={export_format}")

    assert response.status_code == 422
    assert response.headers["content-type"] == "application/json"
    assert "content-disposition" not in response.headers
    assert response.json()["code"] == "export.row_limit_exceeded"
    assert response.json()["params"] == {"max_rows": len(CASE_TITLES) - 1}
    assert response.json()["detail"] == (
        f"Export exceeds the limit of {len(CASE_TITLES) - 1} rows; narrow the filters and retry"
    )
    # A narrower filter under the same limit succeeds.
    narrowed = client.get(f"{EXPORT_URL}?format={export_format}&lifecycle=scheduled")
    assert narrowed.status_code == 200


def test_export_exactly_at_the_row_limit_succeeds(
    postgres_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    seed = _seed(postgres_engine)
    client = _client(postgres_engine, seed.organization_id, seed.lead_id)
    _with_max_rows(monkeypatch, len(CASE_TITLES))

    assert len(_csv_rows(client)) == len(CASE_TITLES)


def test_export_writes_a_structured_log_without_row_content(postgres_engine: Engine) -> None:
    seed = _seed(postgres_engine)
    client = _client(postgres_engine, seed.organization_id, seed.lead_id)
    stream = io.StringIO()
    handler = logging.StreamHandler(stream)
    handler.setFormatter(JsonFormatter())
    previous_level = APP_LOGGER.level
    APP_LOGGER.addHandler(handler)
    APP_LOGGER.setLevel(logging.INFO)
    try:
        assert client.get(f"{EXPORT_URL}?format=xlsx").status_code == 200
    finally:
        APP_LOGGER.removeHandler(handler)
        APP_LOGGER.setLevel(previous_level)

    import json

    [entry] = [
        json.loads(line)
        for line in stream.getvalue().splitlines()
        if '"management_export"' in line
    ]
    assert entry["export_format"] == "xlsx"
    assert entry["row_count"] == len(CASE_TITLES)
    assert entry["organization_id"] == str(seed.organization_id)
    assert entry["actor_user_id"] == str(seed.lead_id)
    assert "dropped_fields" not in entry
    assert "Plain English" not in stream.getvalue()


def test_export_requires_authentication() -> None:
    assert TestClient(create_app()).get(f"{EXPORT_URL}?format=csv").status_code == 401
