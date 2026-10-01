"""CSV / XLSX rendering of a management export snapshot.

Both formats are built from the same row model (`_case_rows`), so they cannot drift from each
other. Every row comes from `ManagementCaseSummary`, i.e. from the same authorized snapshot the
JSON list uses; nothing here queries or recomputes metrics.
"""

import csv
import io
import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import datetime
from enum import StrEnum
from typing import Any
from urllib.parse import quote
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from easyaudit_next.management.schemas import ManagementCaseExportSnapshot, ManagementCaseSummary

CASES_SHEET = "review_cases"
METADATA_SHEET = "metadata"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# A text cell starting with one of these may be evaluated as a formula by spreadsheet software.
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
# Everything outside the XML 1.0 `Char` production (and therefore not representable in XLSX):
# allowed are #x9 | #xA | #xD | [#x20-#xD7FF] | [#xE000-#xFFFD] | [#x10000-#x10FFFF].
_XML_ILLEGAL = re.compile("[^\x09\x0a\x0d\x20-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]")

Cell = str | int | None


class ExportFormat(StrEnum):
    CSV = "csv"
    XLSX = "xlsx"


@dataclass(frozen=True, slots=True)
class ExportDocument:
    content: bytes
    media_type: str
    filename: str

    @property
    def content_disposition(self) -> str:
        return f"attachment; filename=\"{self.filename}\"; filename*=UTF-8''{quote(self.filename)}"


def protect_text(value: str) -> str:
    """Neutralize formula injection and unrepresentable characters in a text cell."""
    cleaned = _XML_ILLEGAL.sub("\ufffd", value)
    if cleaned.startswith(_FORMULA_TRIGGERS):
        return "'" + cleaned
    return cleaned


def _iso(value: datetime | None, tz: ZoneInfo) -> str | None:
    return None if value is None else value.astimezone(tz).isoformat()


# (column name, extractor). Extractors return int for counters, else raw text or None.
_COLUMNS: tuple[tuple[str, Callable[[ManagementCaseSummary, ZoneInfo], Cell]], ...] = (
    ("id", lambda c, tz: str(c.id)),
    ("review_plan_id", lambda c, tz: None if c.review_plan_id is None else str(c.review_plan_id)),
    ("title", lambda c, tz: c.title),
    ("scenario_key", lambda c, tz: c.scenario_key),
    ("scenario_version", lambda c, tz: c.scenario_version),
    ("lifecycle", lambda c, tz: c.lifecycle.value),
    ("planned_start_at", lambda c, tz: _iso(c.planned_start_at, tz)),
    ("planned_end_at", lambda c, tz: _iso(c.planned_end_at, tz)),
    ("deadline_bucket", lambda c, tz: c.deadline_bucket.value),
    ("findings_total", lambda c, tz: c.findings.total),
    ("findings_open", lambda c, tz: c.findings.open),
    ("findings_rectifying", lambda c, tz: c.findings.rectifying),
    ("findings_verifying", lambda c, tz: c.findings.verifying),
    ("findings_closed", lambda c, tz: c.findings.closed),
    ("findings_voided", lambda c, tz: c.findings.voided),
    ("actions_total", lambda c, tz: c.actions.total),
    ("actions_todo", lambda c, tz: c.actions.todo),
    ("actions_in_progress", lambda c, tz: c.actions.in_progress),
    ("actions_done", lambda c, tz: c.actions.done),
    ("actions_cancelled", lambda c, tz: c.actions.cancelled),
    ("actions_overdue", lambda c, tz: c.actions.overdue),
    ("actions_due_soon", lambda c, tz: c.actions.due_soon),
)

EXPORT_COLUMNS: tuple[str, ...] = tuple(name for name, _ in _COLUMNS)


def _cell(value: Cell) -> Cell:
    return protect_text(value) if isinstance(value, str) else value


def _case_rows(snapshot: ManagementCaseExportSnapshot, tz: ZoneInfo) -> list[list[Cell]]:
    return [[_cell(extract(item, tz)) for _, extract in _COLUMNS] for item in snapshot.items]


def _metadata_rows(
    snapshot: ManagementCaseExportSnapshot, generated_at: datetime, tz: ZoneInfo
) -> list[list[Cell]]:
    filters = snapshot.filters
    pairs: list[tuple[str, Cell]] = [
        ("generated_at", _iso(generated_at, tz)),
        ("as_of", _iso(snapshot.as_of, tz)),
        ("timezone", tz.key),
        ("organization_name", snapshot.organization_name),
        (
            "filter_review_plan_id",
            None if filters.review_plan_id is None else str(filters.review_plan_id),
        ),
        ("filter_lifecycle", None if filters.lifecycle is None else filters.lifecycle.value),
        ("filter_deadline_status", filters.deadline_status.value),
        ("row_count", len(snapshot.items)),
    ]
    return [[key, _cell(value)] for key, value in pairs]


def render_export(
    snapshot: ManagementCaseExportSnapshot,
    export_format: ExportFormat,
    *,
    generated_at: datetime,
    tz: ZoneInfo,
) -> ExportDocument:
    cases = _case_rows(snapshot, tz)
    metadata = _metadata_rows(snapshot, generated_at, tz)
    stamp = generated_at.astimezone(tz).strftime("%Y%m%dT%H%M%z")
    filename = f"review-cases-{stamp}.{export_format.value}"
    if export_format is ExportFormat.CSV:
        return ExportDocument(_render_csv(metadata, cases), CSV_MEDIA_TYPE, filename)
    return ExportDocument(_render_xlsx(metadata, cases), XLSX_MEDIA_TYPE, filename)


def _render_csv(metadata: Sequence[Sequence[Cell]], cases: Sequence[Sequence[Cell]]) -> bytes:
    # Metadata first as `key,value` lines, then one blank line, then header and data. A person
    # opening the file sees the context before the data; a parser skips to the first blank line.
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    for row in metadata:
        writer.writerow(["" if value is None else value for value in row])
    writer.writerow([])
    writer.writerow(EXPORT_COLUMNS)
    for row in cases:
        writer.writerow(["" if value is None else value for value in row])
    # BOM so Excel detects UTF-8 and does not garble Chinese text.
    return buffer.getvalue().encode("utf-8-sig")


def _render_xlsx(metadata: Sequence[Sequence[Cell]], cases: Sequence[Sequence[Cell]]) -> bytes:
    workbook = Workbook(write_only=True)
    for title, header, rows in (
        (CASES_SHEET, EXPORT_COLUMNS, cases),
        (METADATA_SHEET, None, metadata),
    ):
        sheet = workbook.create_sheet(title)
        if header is not None:
            sheet.append(list(header))
        for row in rows:
            cells: list[Any] = []
            for value in row:
                if isinstance(value, str):
                    # Explicit string type: openpyxl would otherwise infer a formula from "=".
                    cell = WriteOnlyCell(sheet, value=value)
                    cell.data_type = "s"
                    cells.append(cell)
                else:
                    cells.append(value)
            sheet.append(cells)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
